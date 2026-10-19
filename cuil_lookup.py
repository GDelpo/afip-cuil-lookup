"""Consulta CUIL/CUIT por DNI + sexo en la página pública de ARCA (ex AFIP) "CUIT Digital".

Uso:
    python cuil_lookup.py 12345678 M
    python cuil_lookup.py --input personas.csv
    python cuil_lookup.py --input personas.xlsx --output resultados.xlsx
"""

import argparse
import csv
import io
import json
import sys
import time
from pathlib import Path

from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright

URL_FORM = "https://serviciosweb.afip.gob.ar/publico/cuitOnline/InfoPersonal.aspx"
URL_RESULT = "**/ClaveFiscalQR.aspx"
SEX_OPTIONS = {"M": "Masculino", "F": "Femenino", "X": "No binario"}
TIMEOUT_MS = 10_000
FIELDS = ["dni", "sexo", "nombre", "cuit", "error"]


def lookup(page, dni: str, sex: str, timeout_ms: int = TIMEOUT_MS) -> dict:
    """Run one query on an open page and return a result dict (never raises)."""
    dni = "".join(ch for ch in str(dni) if ch.isdigit())
    sex = str(sex).strip().upper()[:1]
    result = {"dni": dni, "sexo": sex, "nombre": None, "cuit": None, "error": None}
    if not dni or sex not in SEX_OPTIONS:
        result["error"] = "DNI vacío o sexo inválido (usar M, F o X)"
        return result

    try:
        page.goto(URL_FORM, timeout=timeout_ms)
        page.fill("#txtDocumento", dni)
        page.select_option("#ddlGenero", label=SEX_OPTIONS[sex])
        page.click("#btnContinuar")
        page.wait_for_url(URL_RESULT, timeout=timeout_ms)
        result["nombre"] = page.inner_text("#h2Nombre", timeout=timeout_ms).strip()
        result["cuit"] = page.inner_text("#lblCuit", timeout=timeout_ms).strip()
    except PlaywrightError as exc:
        # ARCA muestra el motivo (DNI inexistente, sexo que no coincide) en #divMensajeError
        message = page.locator("#divMensajeError")
        if message.count() and message.inner_text().strip():
            result["error"] = message.inner_text().strip()
        else:
            result["error"] = f"{type(exc).__name__}: {str(exc).splitlines()[0]}"
    return result


def read_people(path: Path) -> list[dict]:
    """Read rows with 'dni' and 'sexo' columns from a .csv or .xlsx file."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip().lower() for h in next(rows)]
        people = [dict(zip(header, row)) for row in rows if any(row)]
    else:
        with path.open(newline="", encoding="utf-8-sig") as fh:
            people = [
                {k.strip().lower(): v for k, v in row.items()}
                for row in csv.DictReader(fh)
            ]
    missing = {"dni", "sexo"} - set(people[0] if people else {})
    if missing:
        sys.exit(f"{path}: faltan columnas {sorted(missing)}")
    return people


def write_results(results: list[dict], output: Path | None) -> None:
    """Write results as JSON (stdout or .json), .csv or .xlsx depending on the extension."""
    if output is None:
        print(json.dumps(results, ensure_ascii=False, indent=2))
        return
    suffix = output.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.worksheets[0]
        ws.append(FIELDS)
        for r in results:
            ws.append([r[f] for f in FIELDS])
        wb.save(output)
    elif suffix == ".csv":
        with output.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(results)
    else:
        output.write_text(
            json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(f"{len(results)} resultados -> {output}", file=sys.stderr)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Consulta CUIL/CUIT por DNI + sexo en ARCA (ex AFIP)."
    )
    parser.add_argument("dni", nargs="?", help="DNI a consultar")
    parser.add_argument("sexo", nargs="?", help="M, F o X")
    parser.add_argument("--input", type=Path, help="CSV o XLSX con columnas dni,sexo")
    parser.add_argument(
        "--output",
        type=Path,
        help="Archivo .json, .csv o .xlsx (default: JSON por stdout)",
    )
    parser.add_argument(
        "--delay", type=float, default=1.0, help="Segundos entre consultas (default: 1)"
    )
    parser.add_argument("--headed", action="store_true", help="Mostrar el navegador")
    args = parser.parse_args()
    # La consola de Windows usa cp1252 por defecto y rompe los acentos del JSON
    for stream in (sys.stdout, sys.stderr):
        if isinstance(stream, io.TextIOWrapper):
            stream.reconfigure(encoding="utf-8")

    if args.input:
        people = read_people(args.input)
    elif args.dni and args.sexo:
        people = [{"dni": args.dni, "sexo": args.sexo}]
    else:
        parser.error("pasá DNI y sexo, o --input")

    results = []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not args.headed)
        page = browser.new_page()
        for i, person in enumerate(people):
            if i:
                time.sleep(args.delay)
            results.append(lookup(page, person["dni"], person["sexo"]))
        browser.close()

    write_results(results, args.output)
    if any(r["error"] for r in results):
        sys.exit(1)


if __name__ == "__main__":
    main()
